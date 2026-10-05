#!/usr/bin/env python3
"""Build Protocol-v2 / experiment-suite aligned Excel workbook for reviewers."""
from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
CSV = ROOT / "data" / "csv"
OUT = ROOT / "data" / "AI_NET_Experimental_Dataset_SuiteAligned.xlsx"

# Frozen hidden-test counts in the public experiment_suite evaluators
SUITE_TEST_COUNTS = {
    "TP01": 10,
    "TP02": 20,
    "TP03": 20,
    "TP04": 20,
    "TP05": 20,
    "TP06": 20,
    "TP07": 20,
    "TP08": 20,
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SECTION_FILL = PatternFill("solid", fgColor="D6EAF8")


def read_csv(name: str) -> list[dict[str, str]]:
    with (CSV / name).open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def autosize(ws, max_width: int = 28) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min(max_width, max(len(str(c.value or "")) for c in col) + 2)
        ws.column_dimensions[letter].width = max(10, width)


def write_sheet(wb: Workbook, title: str, rows: list[dict], columns: list[str] | None = None) -> None:
    ws = wb.create_sheet(title)
    if not rows:
        ws.append(["(empty)"])
        return
    cols = columns or list(rows[0].keys())
    ws.append(cols)
    for row in rows:
        ws.append([coerce(row.get(c, "")) for c in cols])
    style_header(ws, len(cols))
    autosize(ws)


def coerce(value):
    if value is None or value == "":
        return value
    if isinstance(value, (int, float, bool)):
        return value
    if value in ("True", "False"):
        return value == "True"
    try:
        if "." in value:
            return float(value)
        return int(value)
    except ValueError:
        return value


def descriptive(observations: list[dict]) -> list[dict]:
    metrics = [
        ("completion_time_min", "CompletionTime_min"),
        ("test_failures", "TestFailures"),
        ("code_smells", "CodeSmells"),
        ("cyclomatic_complexity", "CyclomaticComplexity"),
        ("maintainability_index", "MaintainabilityIndex"),
        ("loc", "LOC"),
        ("nasa_tlx", "NASA_TLX"),
    ]
    out = []
    for key, label in metrics:
        for condition in ("Traditional", "AI"):
            vals = [float(r[key]) for r in observations if r["condition"] == condition]
            n = len(vals)
            mean = sum(vals) / n
            var = sum((v - mean) ** 2 for v in vals) / (n - 1)
            sd = var ** 0.5
            se = sd / (n ** 0.5)
            out.append(
                {
                    "Metric": label,
                    "Condition": condition,
                    "N": n,
                    "Mean": round(mean, 4),
                    "SD": round(sd, 4),
                    "SE": round(se, 4),
                    "CI95_Lower": round(mean - 1.96 * se, 4),
                    "CI95_Upper": round(mean + 1.96 * se, 4),
                }
            )
    return out


def task_level(observations: list[dict]) -> list[dict]:
    groups = defaultdict(list)
    for r in observations:
        groups[(r["task_pair"], r["condition"])].append(r)
    out = []
    for (task, condition), rows in sorted(groups.items()):
        times = [float(x["completion_time_min"]) for x in rows]
        fails = [float(x["test_failures"]) for x in rows]
        out.append(
            {
                "task_pair": task,
                "condition": condition,
                "N": len(rows),
                "mean_completion_time_min": round(sum(times) / len(times), 4),
                "mean_test_failures": round(sum(fails) / len(fails), 4),
            }
        )
    return out


def suite_inventory(task_pairs: list[dict]) -> list[dict]:
    out = []
    for row in task_pairs:
        tp = row["task_pair"]
        out.append(
            {
                "task_pair": tp,
                "variant": row["variant"],
                "task_name": row["task_name"],
                "capability": {
                    "TP01": "REST CRUD + validation",
                    "TP02": "Authorization + business rules",
                    "TP03": "Query + aggregation",
                    "TP04": "State-transition workflow",
                    "TP05": "External API integration",
                    "TP06": "File import + validation",
                    "TP07": "Refactoring + defect correction",
                    "TP08": "Reporting + date logic",
                }[tp],
                "suite_hidden_tests_frozen": SUITE_TEST_COUNTS[tp],
                "pilot_difficulty_rating_1_5": row["pilot_difficulty_rating_1_5"],
                "contracts_path": f"Tasks/{tp}/{row['variant']}/Contracts",
                "participant_path": f"Tasks/{tp}/{row['variant']}/Participant",
                "evaluation_path": f"Tasks/{tp}/{row['variant']}/Evaluation",
                "task_spec_glob": f"task_specs/{tp}_{row['variant']}_*.md",
            }
        )
    return out


def readme_rows() -> list[list[str]]:
    return [
        ["Field", "Value"],
        ["Workbook", "AI_NET_Experimental_Dataset_SuiteAligned.xlsx"],
        ["Purpose", "Reviewer-facing dataset aligned to experiment_suite + Protocol v2.0"],
        ["Replaces", "AI_NET_Experimental_Dataset_Revised.xlsx (older T1–T8 flat schema)"],
        ["Observations", "960 = 60 participants × 8 task pairs × 2 conditions"],
        ["Task IDs", "TP01–TP08 with matched A/B variants"],
        ["Design", "Counterbalanced crossover sequences A–D; period 1/2"],
        ["Primary productivity", "completion_time_min / completion_time_sec"],
        ["Primary correctness", "test_failures (hidden suite)"],
        ["Structural quality", "code_smells, cyclomatic_complexity, maintainability_index"],
        ["Workload", "nasa_tlx"],
        ["Analyzer package", "SonarAnalyzer.CSharp 10.35.0.4138"],
        ["Analyzer config hash", "c84cde8e4805433f (see analyzer/CONFIG_HASH.txt)"],
        ["Metrics package", "Microsoft.CodeAnalysis.Metrics 5.6.0"],
        ["CSV sources", "data/csv/*.csv"],
        ["Code & tasks", "https://github.com/MaazTariq409/experiment_suite"],
    ]


def main() -> None:
    observations = read_csv("observations_master.csv")
    participants = read_csv("participants_baseline.csv")
    sessions = read_csv("sessions_summary.csv")
    task_pairs = read_csv("task_pairs_master.csv")

    # Enrich observations with experience_tier from participants
    tier = {p["participant_id"]: p["experience_tier"] for p in participants}
    for row in observations:
        row["experience_tier"] = tier.get(row["participant_id"], "")
        row["suite_hidden_tests_frozen"] = SUITE_TEST_COUNTS[row["task_pair"]]

    obs_cols = [
        "observation_id",
        "participant_id",
        "experience_tier",
        "sequence",
        "period",
        "condition",
        "task_pair",
        "task_variant",
        "task_name",
        "task_order_in_session",
        "completion_time_sec",
        "completion_time_min",
        "time_limit_reached",
        "hidden_tests_total",
        "suite_hidden_tests_frozen",
        "tests_passed",
        "test_failures",
        "code_smells",
        "cyclomatic_complexity",
        "maintainability_index",
        "loc",
        "nasa_tlx",
        "programming_years",
        "ai_familiarity_score",
        "protocol_deviation",
        "environment_version",
        "test_suite_version",
        "analyzer_version",
        "data_provenance",
    ]

    wb = Workbook()
    # README first
    ws = wb.active
    ws.title = "README"
    for r in readme_rows():
        ws.append(r)
    ws["A1"].fill = HEADER_FILL
    ws["B1"].fill = HEADER_FILL
    ws["A1"].font = HEADER_FONT
    ws["B1"].font = HEADER_FONT
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=1):
        for cell in row:
            cell.fill = SECTION_FILL
            cell.font = Font(bold=True)
    autosize(ws, 40)
    ws.column_dimensions["B"].width = 90

    write_sheet(wb, "Observations", observations, obs_cols)
    write_sheet(wb, "Participants", participants)
    write_sheet(wb, "Sessions", sessions)
    write_sheet(wb, "TaskPairs_Suite", suite_inventory(task_pairs))
    write_sheet(wb, "Descriptive_By_Condition", descriptive(observations))
    write_sheet(wb, "TaskLevel_Means", task_level(observations))

    # Mapping from legacy Excel
    mapping = [
        {"legacy_TaskID": "T1", "task_pair": "TP01", "capability": "REST CRUD + validation"},
        {"legacy_TaskID": "T2", "task_pair": "TP02", "capability": "Authorization + business rules"},
        {"legacy_TaskID": "T3", "task_pair": "TP03", "capability": "Query + aggregation"},
        {"legacy_TaskID": "T4", "task_pair": "TP04", "capability": "State-transition workflow"},
        {"legacy_TaskID": "T5", "task_pair": "TP05", "capability": "External API integration"},
        {"legacy_TaskID": "T6", "task_pair": "TP06", "capability": "File import + validation"},
        {"legacy_TaskID": "T7", "task_pair": "TP07", "capability": "Refactoring + defect correction"},
        {"legacy_TaskID": "T8", "task_pair": "TP08", "capability": "Reporting + date logic"},
    ]
    write_sheet(wb, "LegacyExcel_Mapping", mapping)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Observations: {len(observations)}")


if __name__ == "__main__":
    main()
